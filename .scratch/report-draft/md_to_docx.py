"""Render a report Markdown file into the official FinTechathon Word template.

Usage: python md_to_docx.py <template.docx> <report.md> <out.docx> <title> <subtitle>
Relative links become GitHub URLs on main; images are embedded.
"""
import posixpath
import re
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

sys.stdout.reconfigure(encoding="utf-8")
REPO_URL = "https://github.com/DrErwin/hackerthon-bank-agent/blob/main/"
TEAM, TRACK, DATE, PRODUCT = "KOSX1", "国内赛道 · 人工智能赛道（AI Banking Agent）", "2026-11-01", "SecurityFlex"
CJK_FONT = "Microsoft YaHei"
LINK_BLUE = RGBColor(0x31, 0x5F, 0xD8)
SKIP_PREFIX = ("- 作品名", "- 队伍名", "- 赛道", "- 提交日期", "赛道：", "日期：", f"{PRODUCT}｜")

template, md_path, out_path, title, subtitle = sys.argv[1:6]
md_path = Path(md_path)
repo_root = Path(md_path).resolve().parents[2]
md_dir_rel = md_path.resolve().parent.relative_to(repo_root).as_posix()

doc = docx.Document(template)
body = doc.element.body

# ---- cover: keep the brand line, retitle, fill the info table, drop the skeleton
paras = doc.paragraphs
paras[1].runs[0].text = title
paras[2].runs[0].text = subtitle
table = doc.tables[0]
for row, value in zip(table.rows, (TEAM, TRACK, DATE)):
    cell = row.cells[1].paragraphs[0]
    cell.runs[0].text = value
    cell.runs[0].font.color.rgb = RGBColor(0x17, 0x21, 0x3A)
product_row = table.rows[0]._tr.__copy__()
table._tbl.insert(table._tbl.index(table.rows[0]._tr), product_row)
first = table.rows[0]
first.cells[0].paragraphs[0].runs[0].text = "作品名称 / Project"
first.cells[1].paragraphs[0].runs[0].text = PRODUCT
for p in paras[3:]:
    body.remove(p._p)
doc.add_paragraph()

# The template names "Arial Unicode MS" (rarely installed; substitutes render Latin
# as Greek). Pin every style, run and the theme to one widely available font.
def pin_fonts(rpr_parent):
    for rfonts in rpr_parent.iter(qn("w:rFonts")):
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        if rfonts.get(qn("w:ascii")) not in ("Consolas", "Courier"):
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), CJK_FONT)


for style in doc.styles:
    rpr = style.element.get_or_add_rPr() if style.type == 1 else None
    if rpr is not None:
        rpr.get_or_add_rFonts()
pin_fonts(doc.styles.element)
pin_fonts(body)
# The footer is a note to whoever fills the template, not report content.
for section in doc.sections:
    for p in section.footer.paragraphs:
        for run in p.runs:
            run.text = ""
    section.footer.paragraphs[0].add_run(f"{PRODUCT} · {TEAM}").font.size = Pt(8)


def resolve(target: str) -> str:
    if re.match(r"^[a-z]+:", target):
        return target
    path, _, anchor = target.partition("#")
    full = posixpath.normpath(posixpath.join(md_dir_rel, path))
    return REPO_URL + full + (f"#{anchor}" if anchor else "")


def add_hyperlink(paragraph, text, url, bold=False, size=None):
    rel = paragraph.part.relate_to(url, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), str(LINK_BLUE))
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    if bold:
        rpr.append(OxmlElement("w:b"))
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text, size=None, bold=False):
    text = text.replace("\\*", "*").replace("\\|", "|")
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            add_inline(paragraph, piece[2:-2], size, True)
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = size or Pt(9.5)
            run.bold = bold
        elif piece.startswith("["):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", piece).groups()
            add_hyperlink(paragraph, label.strip("`"), resolve(url), bold, size)
        else:
            run = paragraph.add_run(piece)
            run.bold = bold or None
            if size:
                run.font.size = size


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C9CFDA")
        borders.append(el)
    tc_pr.append(borders)


def shade(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def split_row(line):
    cells = re.split(r"(?<!\\)\|", line.strip().strip("|"))
    return [c.strip() for c in cells]


def add_table(rows):
    header, data = rows[0], rows[2:]
    t = doc.add_table(rows=1 + len(data), cols=len(header))
    t.autofit = True
    for r, values in enumerate([header] + data):
        for c in range(len(header)):
            cell = t.cell(r, c)
            set_cell_border(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, values[c] if c < len(values) else "", Pt(9), bold=(r == 0))
            if r == 0:
                shade(cell, "EEF2FA")
    doc.add_paragraph()


def add_image(alt, target):
    path = (md_path.parent / target).resolve()
    with Image.open(path) as im:
        w, h = im.size
    if h > w:  # phone screenshot: cap height and width
        height = min(Inches(7.0), Inches(3.2) * h / w)
        width = int(height * w / h)
    else:
        width = Inches(6.3)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=int(width))


def add_caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_inline(p, text, Pt(9))
    for run in p.runs:
        run.font.color.rgb = RGBColor(0x61, 0x6C, 0x7F)


lines = md_path.read_text(encoding="utf-8-sig").splitlines()
started_body = False
i = 0
previous_image = False
while i < len(lines):
    line = lines[i].rstrip()
    stripped = line.strip()
    i += 1
    if not stripped:
        continue
    if stripped.startswith("# ") and not started_body:
        continue
    if not started_body and stripped.startswith(SKIP_PREFIX):
        continue
    if stripped.startswith("|"):
        rows = [split_row(stripped)]
        while i < len(lines) and lines[i].strip().startswith("|"):
            rows.append(split_row(lines[i]))
            i += 1
        add_table(rows)
        previous_image = False
        continue
    heading = re.match(r"^(#{2,4}) (.+)$", stripped)
    if heading:
        started_body = True
        level = len(heading.group(1)) - 1
        doc.add_paragraph(heading.group(2), style=f"Heading {level}")
        previous_image = False
        continue
    image = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", stripped)
    if image:
        add_image(*image.groups())
        previous_image = True
        continue
    if previous_image and re.match(r"^图", stripped):
        add_caption(stripped)
        previous_image = False
        continue
    previous_image = False
    bullet = re.match(r"^(\s*)[-*] (.+)$", line)
    number = re.match(r"^(\s*)\d+\. (.+)$", line)
    if bullet:
        depth = len(bullet.group(1)) // 2
        style = "List Bullet" if depth == 0 else "List Bullet 2"
        add_inline(doc.add_paragraph(style=style), bullet.group(2))
    elif number:
        # Keep the literal number: Word's List Number would chain across the document.
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        add_inline(p, stripped)
    elif stripped.startswith("> "):
        p = doc.add_paragraph(style="Quote")
        add_inline(p, stripped[2:])
    else:
        add_inline(doc.add_paragraph(), stripped)

for run in body.iter(qn("w:r")):
    rpr = run.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        run.insert(0, rpr)
    if rpr.find(qn("w:rFonts")) is None:
        rpr.insert(0, OxmlElement("w:rFonts"))
pin_fonts(body)
doc.save(out_path)
print("saved", out_path)
